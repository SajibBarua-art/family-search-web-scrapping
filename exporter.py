"""
Excel report exporter module using pandas and openpyxl.
Formats batch outputs into clean, timestamped Excel workbooks.
"""
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import logging

from config import REPORTS_DIR

logger = logging.getLogger(__name__)


def format_excel_sheet(file_path: Path) -> None:
    """
    Applies professional styling to the generated Excel workbook using openpyxl:
    - Bold white text on navy header banner
    - Thin border gridlines
    - Auto-adjusted column widths
    - Freezes the header row
    """
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    
    # Palette & Styles
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    data_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    
    # Freeze header row
    ws.freeze_panes = "A2"
    
    # Style Header Row
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    
    # Style Data Rows
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.font = data_font
            cell.border = thin_border
            if isinstance(cell.value, (int, float)):
                cell.alignment = center_align
            else:
                cell.alignment = left_align
                
    # Auto-fit Column Widths with padding
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        # Cap max column width to 50 for readability, min width 12
        ws.column_dimensions[col_letter].width = max(min(max_len + 4, 50), 12)
        
    wb.save(file_path)


def export_batch_report(
    batch_number: int,
    batch_records: List[Dict[str, Any]],
    reports_dir: Path = REPORTS_DIR,
) -> Path:
    """
    Generate an Excel report (.xlsx) for a processed batch of search queries.
    File naming convention:
        reports/familysearch_batch_{batch_number}_{YYYYMMDD_HHMMSS}.xlsx
    """
    reports_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"familysearch_batch_{batch_number}_{timestamp}.xlsx"
    file_path = reports_dir / filename
    
    rows: List[Dict[str, Any]] = []
    
    for item in batch_records:
        query_info = item.get("query", {})
        locations = item.get("locations", [])
        is_cached = item.get("is_cached", False)
        
        search_qp = query_info.get("qp", "")
        search_lat = query_info.get("lat")
        search_lng = query_info.get("lng")
        search_radius = query_info.get("radius")
        query_sig = query_info.get("query_signature", "")
        
        if not locations:
            # Query returned 0 locations, include an explicit row for completeness
            rows.append({
                "Batch Number": batch_number,
                "Query Signature": query_sig,
                "Search QP": search_qp,
                "Search Lat": search_lat,
                "Search Lng": search_lng,
                "Search Radius (mi)": search_radius,
                "Source Status": "CACHED (0 Results)" if is_cached else "FETCHED (0 Results)",
                "Location Entity ID": None,
                "Location Name": "No locations found within radius",
                "Location Type": None,
                "Street Address": None,
                "City": None,
                "State/Region": None,
                "Postal Code": None,
                "Country": None,
                "Phone": None,
                "Website URL": None,
                "Result Lat": None,
                "Result Lng": None,
                "Distance (Miles)": None,
                "Hours / Notes": None,
            })
        else:
            for loc in locations:
                rows.append({
                    "Batch Number": batch_number,
                    "Query Signature": query_sig,
                    "Search QP": search_qp,
                    "Search Lat": search_lat,
                    "Search Lng": search_lng,
                    "Search Radius (mi)": search_radius,
                    "Source Status": "CACHED" if is_cached else "FETCHED",
                    "Location Entity ID": loc.get("location_entity_id"),
                    "Location Name": loc.get("location_name"),
                    "Location Type": loc.get("location_type"),
                    "Street Address": loc.get("street_address"),
                    "City": loc.get("city"),
                    "State/Region": loc.get("state_region"),
                    "Postal Code": loc.get("postal_code"),
                    "Country": loc.get("country"),
                    "Phone": loc.get("phone"),
                    "Website URL": loc.get("website_url"),
                    "Result Lat": loc.get("res_lat"),
                    "Result Lng": loc.get("res_lng"),
                    "Distance (Miles)": loc.get("distance_miles"),
                    "Hours / Notes": loc.get("hours_or_notes"),
                })
                
    df = pd.DataFrame(rows)
    
    # Save to Excel
    df.to_excel(file_path, index=False, engine="openpyxl")
    
    # Apply professional openpyxl styling
    format_excel_sheet(file_path)
    
    logger.info("Exported batch %d report with %d rows to: %s", batch_number, len(rows), file_path)
    return file_path
